from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from CSVs_to_excel import combine_csv_files_to_excel


def test_colliding_long_filenames_keep_distinct_worksheets(tmp_path: Path):
    prefix = "a_very_long_csv_filename_prefix_"
    first = prefix + "alpha.csv"
    second = prefix + "beta.csv"

    pd.DataFrame({"value": [1]}).to_csv(tmp_path / first, index=False)
    pd.DataFrame({"value": [2]}).to_csv(tmp_path / second, index=False)

    combine_csv_files_to_excel(tmp_path)

    workbook = load_workbook(tmp_path / "combined_data.xlsx", read_only=True)
    base_name = Path(first).stem[:31]

    assert workbook.sheetnames == [base_name, f"{base_name[:29]}_2"]
    assert list(workbook[base_name].values) == [("value",), (1,)]
    assert list(workbook[f"{base_name[:29]}_2"].values) == [("value",), (2,)]


def test_non_colliding_worksheet_names_keep_existing_behavior(tmp_path: Path):
    pd.DataFrame({"value": [3]}).to_csv(tmp_path / "ordinary_name.csv", index=False)
    pd.DataFrame({"value": [4]}).to_csv(tmp_path / "name[with]invalid.csv", index=False)

    combine_csv_files_to_excel(tmp_path)

    workbook = load_workbook(tmp_path / "combined_data.xlsx", read_only=True)

    assert workbook.sheetnames == ["name_with_invalid", "ordinary_name"]
